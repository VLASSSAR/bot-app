from collections import Counter
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import COMPANIES
from .models import Role


def hyperlink(paragraph, label, url):
    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1557A0")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def build_report(result, timezone, sample=False):
    doc = Document()
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [("Title", 24), ("Heading 1", 15), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
    doc.add_paragraph("Вакансии российского бигтеха", "Title")
    if sample:
        doc.add_paragraph("Пример отчета — сокращенная подборка", "Subtitle")
        doc.styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph(
        "Продуктовый менеджмент, управление проектами и технический продуктовый менеджмент."
    )
    stamp = result.collected_at.astimezone(timezone).strftime("%d.%m.%Y %H:%M")
    doc.add_paragraph(
        f"Данные собраны {stamp} ({timezone.key}). В подборке {len(result.vacancies)} вакансий."
    )
    companies = ", ".join(COMPANIES[k][0] for k in result.filters.companies)
    doc.add_paragraph(
        f"Компании: {companies}. Период: {result.filters.days} дней. "
        f"Формат: {'только удалённо' if result.filters.remote_only else 'все варианты'}."
    )
    doc.add_paragraph(
        "Приоритет: сайты работодателей, Getmatch и Хабр Карьера. "
        "Вакансии без даты публикации отобраны по дате первого обнаружения "
        "и вынесены отдельно. При первом запуске они могут оказаться старше выбранного периода. "
        "Неуказанная география отмечена в карточке; работу из России нужно уточнить."
    )
    counts = Counter(v.role for v in result.vacancies)
    doc.add_paragraph("По ролям", "Heading 1")
    for role in Role:
        doc.add_paragraph(f"{role}: {counts[role]}")
    doc.add_paragraph("Проверка источников", "Heading 1")
    labels = {"ok": "Проверен", "partial": "Частично", "error": "Недоступен"}
    for status in result.sources:
        p = doc.add_paragraph()
        p.add_run(f"{status.name} — {labels[status.state]}. ").bold = True
        if status.message:
            p.add_run(status.message + " ")
        hyperlink(p, "Открыть источник", status.url)
    doc.add_paragraph(
        "Охват ограничен публичной выдачей доступных источников и выбранными "
        "профилями компаний. Отчет не гарантирует наличие всех вакансий группы компаний. "
        "Ссылки могут закрыться после сбора."
    )
    if not result.vacancies:
        doc.add_paragraph("Подходящие вакансии не найдены", "Heading 1")
        doc.add_paragraph("Проверьте статусы источников и расширьте фильтры.")
    for dated, heading in [
        (True, "Вакансии с датой публикации"),
        (False, "Вакансии без даты публикации"),
    ]:
        jobs = [v for v in result.vacancies if (v.published_at is not None) == dated]
        if not jobs:
            continue
        doc.add_page_break()
        doc.add_paragraph(heading, "Heading 1")
        for v in jobs:
            doc.add_paragraph(v.title, "Heading 2")
            p = doc.add_paragraph()
            p.add_run(f"{v.company} · {v.role}").bold = True
            doc.add_paragraph(f"{v.city} · {v.work_format} · Опыт: {v.experience}")
            doc.add_paragraph(f"Зарплата: {v.salary}")
            if not v.region_confirmed:
                doc.add_paragraph("Возможность работы из России не подтверждена источником.")
            date_text = (
                f"Опубликована: {v.published_at.strftime('%d.%m.%Y')}"
                if v.published_at
                else "Дата публикации не указана"
            )
            if v.first_seen:
                date_text += f" · Впервые найдена: {v.first_seen.astimezone(timezone):%d.%m.%Y}"
            doc.add_paragraph(date_text)
            p = doc.add_paragraph()
            hyperlink(p, f"Открыть вакансию — {v.source}", v.url)
            for index, url in enumerate(v.alternate_urls, 1):
                p.add_run(" · ")
                hyperlink(p, f"Другой источник {index}", url)
            # Keep a card together where it fits; long cards can flow naturally.
            paragraphs = doc.paragraphs
            for para in paragraphs[-(7 if not v.region_confirmed else 6) : -1]:
                para.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(10)
    doc.core_properties.title = "Вакансии российского бигтеха"
    doc.core_properties.author = "bot-app"
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
